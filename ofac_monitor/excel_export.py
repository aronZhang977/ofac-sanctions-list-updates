from __future__ import annotations

import logging
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .config import COLUMN_LABELS, enabled_columns, resolve_path
from .dates import address_matches_keywords, format_dob
from .parse_notice import Record

logger = logging.getLogger("ofac_monitor")


def _prepare_row(data: Dict[str, str]) -> Dict[str, str]:
    out = dict(data)
    if "dob" in out:
        out["dob"] = format_dob(out.get("dob") or "")
    return out


def _should_highlight(data: Dict[str, str], cfg: Dict[str, Any]) -> bool:
    excel = cfg.get("excel") or {}
    hl = excel.get("highlight_china_hk") or {}
    if hl.get("enabled", True) is False:
        return False
    keywords = hl.get("keywords") or ["china", "hong kong", "hongkong", "hong-kong"]
    # 以地址为主，国籍/出生地含 China / Hong Kong 也高亮，避免地址为空漏标
    haystack = " ".join(
        [
            data.get("address") or "",
            data.get("nationality") or "",
            data.get("pob") or "",
        ]
    )
    return address_matches_keywords(haystack, keywords)


def cleanup_output_dir(out_dir: Path, cfg: Dict[str, Any]) -> None:
    """Avoid unbounded accumulation of Excel files."""
    excel = cfg.get("excel") or {}
    max_files = int(excel.get("max_files") or 30)
    keep_days = int(excel.get("keep_days") or 0)

    files = sorted(
        [p for p in out_dir.glob("ofac_*.xlsx") if p.is_file()],
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not files:
        return

    keep: set[Path] = set()
    if keep_days > 0:
        cutoff = time.time() - keep_days * 86400
        for p in files:
            if p.stat().st_mtime >= cutoff:
                keep.add(p)
    else:
        keep.update(files[: max(1, max_files)])

    # Always honor max_files even when keep_days is set
    if max_files > 0:
        by_mtime = sorted(keep, key=lambda p: p.stat().st_mtime, reverse=True)
        keep = set(by_mtime[:max_files])

    removed = 0
    for p in files:
        if p not in keep:
            try:
                p.unlink()
                removed += 1
            except OSError as exc:
                logger.warning("删除旧文件失败 %s: %s", p, exc)
    if removed:
        logger.info("已清理旧 Excel %s 个，当前保留 %s 个", removed, len(keep))


def export_records(
    records: Sequence[Record],
    cfg: Dict[str, Any],
    *,
    date_key: str,
    title_slug: str = "ofac_update",
) -> Path:
    columns = enabled_columns(cfg)
    excel = cfg.get("excel") or {}
    out_dir = resolve_path(cfg, excel.get("output_dir") or "output")
    out_dir.mkdir(parents=True, exist_ok=True)

    overwrite = bool(excel.get("overwrite_by_date", True))
    if overwrite:
        filename = f"ofac_{date_key}.xlsx"
    else:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_slug = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in title_slug)[:40]
        filename = f"ofac_{date_key}_{safe_slug}_{stamp}.xlsx"
    path = out_dir / filename

    hl = excel.get("highlight_china_hk") or {}
    fill_color = (hl.get("fill_color") or "FFF2CC").lstrip("#")
    highlight_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

    wb = Workbook()
    ws = wb.active
    ws.title = "OFAC Updates"

    headers = [COLUMN_LABELS.get(c, c) for c in columns]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    highlight_count = 0
    for rec in records:
        data = _prepare_row(rec.as_dict())
        ws.append([data.get(c, "") for c in columns])
        row_idx = ws.max_row
        mark = _should_highlight(data, cfg)
        for cell in ws[row_idx]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if mark:
                cell.fill = highlight_fill
        if mark:
            highlight_count += 1

    for col in ws.columns:
        letter = col[0].column_letter
        max_len = 0
        for cell in col[:200]:
            val = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, min(len(val), 60))
        ws.column_dimensions[letter].width = max(12, max_len + 2)

    wb.save(path)
    if highlight_count:
        logger.info("已高亮 China/Hong Kong 相关行 %s 条", highlight_count)

    cleanup_output_dir(out_dir, cfg)
    return path
